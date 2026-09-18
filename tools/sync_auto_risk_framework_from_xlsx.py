#!/usr/bin/env python3
"""Sync auto-risk-framework HTML from the discussion Excel parameter workbook.

Usage:
  python3 tools/sync_auto_risk_framework_from_xlsx.py \
    [--xlsx data/reports/auto-risk-framework/风险框架参数表_讨论用.xlsx] \
    [--html data/reports/auto-risk-framework/index.html]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = BASE / "data/reports/auto-risk-framework/风险框架参数表_讨论用.xlsx"
DEFAULT_HTML = BASE / "data/reports/auto-risk-framework/index.html"
DEFAULT_HTML_COPY = BASE / "data/reports/auto-risk-framework/auto_risk_framework.html"

DIM_KEYS = ("person", "vehicle", "env", "adas")
DIM_NAME = {"person": "人", "vehicle": "车", "env": "环境", "adas": "智驾"}


def _yn(v) -> bool:
    return str(v or "Y").strip().upper() in {"Y", "YES", "TRUE", "1", "是"}


def _js_str(s: str) -> str:
    return json.dumps(str(s or ""), ensure_ascii=False)


def read_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)

    # meta
    meta = {}
    ws = wb["07_元信息"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        meta[str(row[0]).strip()] = "" if row[2] is None else str(row[2]).strip()

    # weights + labels
    ws = wb["01_风险权重"]
    segments: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[2] is None:
            continue
        sid = str(row[0]).strip()
        if sid.startswith("说明") or "勿改" in sid:
            continue
        try:
            year = int(row[2])
        except (TypeError, ValueError):
            continue
        label = str(row[1] or sid).strip()
        vals = [row[3], row[4], row[5], row[6]]
        if any(v is None for v in vals):
            raise ValueError(f"权重缺数: {sid} {year}")
        nums = [int(round(float(v))) for v in vals]
        total = sum(nums)
        if total != 100:
            raise ValueError(f"权重合计必须=100: {sid} {year} → {total} ({nums})")
        seg = segments.setdefault(sid, {"id": sid, "label": label, "weights": {}, "judge": {}})
        seg["label"] = label
        seg["weights"][year] = dict(zip(DIM_KEYS, nums))

    # judges
    ws = wb["02_判断说明"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[2] is None:
            continue
        sid = str(row[0]).strip()
        try:
            year = int(row[2])
        except (TypeError, ValueError):
            continue
        text = str(row[3] or "").strip()
        if sid not in segments:
            raise ValueError(f"判断说明中有未知分层ID: {sid}")
        if row[1]:
            segments[sid]["label"] = str(row[1]).strip()
        segments[sid]["judge"][year] = text

    # factors
    ws = wb["03_因子布局"]
    factors: dict[str, list] = {k: [] for k in DIM_KEYS}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[3] is None:
            continue
        dim = str(row[0]).strip()
        if dim not in factors:
            continue
        if not _yn(row[9] if len(row) > 9 else "Y"):
            continue
        factors[dim].append(
            {
                "title": str(row[3] or "").strip(),
                "explained": int(round(float(row[4] or 0))),
                "growth": int(round(float(row[5] or 0))),
                "type": str(row[6] or "").strip(),
                "desc": str(row[7] or "").strip(),
                "detail": str(row[8] or "").strip(),
                "sort": int(row[2] or 0),
            }
        )
    for dim in factors:
        factors[dim].sort(key=lambda x: x["sort"])

    # traditional / adas summary tables
    trad = []
    ws = wb["04_传统因子总表"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not _yn(row[4]):
            continue
        trad.append(
            {
                "cat": str(row[0]).strip(),
                "factors": str(row[1] or "").strip(),
                "explain": str(row[2] or "").strip(),
                "status": str(row[3] or "").strip(),
            }
        )

    adas = []
    ws = wb["05_智驾因子总表"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not _yn(row[4]):
            continue
        adas.append(
            {
                "dir": str(row[0]).strip(),
                "obs": str(row[1] or "").strip(),
                "risk": str(row[2] or "").strip(),
                "maturity": str(row[3] or "").strip(),
                "priority": int(row[5] or 99),
            }
        )
    adas.sort(key=lambda x: x["priority"])

    todos = []
    ws = wb["06_待拍板"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        todos.append(
            {
                "n": int(row[0]),
                "topic": str(row[1] or "").strip(),
                "default": str(row[2] or "").strip(),
                "conclusion": str(row[3] or "").strip(),
                "status": str(row[4] or "待讨论").strip(),
            }
        )
    todos.sort(key=lambda x: x["n"])

    # keep segment order as in weight sheet appearance
    order = []
    ws = wb["01_风险权重"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or row[2] is None:
            continue
        sid = str(row[0]).strip()
        if sid in {"fleet", "l2", "ice", "nev"} or (
            sid and sid.replace("_", "").isalnum() and sid[0].isalpha()
        ):
            if sid not in order and sid in segments:
                order.append(sid)

    return {
        "meta": meta,
        "segments": [segments[sid] for sid in order],
        "factors": factors,
        "trad": trad,
        "adas": adas,
        "todos": todos,
    }


def emit_segments_js(segments: list[dict]) -> str:
    blocks = []
    for seg in segments:
        years = sorted(seg["weights"].keys())
        w_lines = []
        for y in years:
            w = seg["weights"][y]
            w_lines.append(
                f"        {y}: {{ person: {w['person']}, vehicle: {w['vehicle']}, env: {w['env']}, adas: {w['adas']} }}"
            )
        j_years = sorted(seg["judge"].keys())
        j_lines = [f"        {y}: {_js_str(seg['judge'].get(y, ''))}" for y in j_years]
        blocks.append(
            "    {\n"
            f"      id: {_js_str(seg['id'])},\n"
            f"      label: {_js_str(seg['label'])},\n"
            "      judge: {\n"
            + ",\n".join(j_lines)
            + "\n      },\n"
            "      weights: {\n"
            + ",\n".join(w_lines)
            + "\n      }\n"
            "    }"
        )
    return "  const SEGMENTS = [\n" + ",\n".join(blocks) + "\n  ];"


def emit_factors_js(factors: dict) -> str:
    parts = []
    for dim in DIM_KEYS:
        items = []
        for f in factors.get(dim, []):
            items.append(
                "      {\n"
                f"        title: {_js_str(f['title'])},\n"
                f"        explained: {f['explained']},\n"
                f"        growth: {f['growth']},\n"
                f"        type: {_js_str(f['type'])},\n"
                f"        desc: {_js_str(f['desc'])},\n"
                f"        detail: {_js_str(f['detail'])}\n"
                "      }"
            )
        parts.append(f"    {dim}: [\n" + ",\n".join(items) + "\n    ]")
    return "  const FACTORS = {\n" + ",\n".join(parts) + "\n  };"


def replace_trad_table(html: str, trad: list[dict]) -> str:
    rows = []
    for t in trad:
        rows.append(
            "<tr>\n"
            f"          <td>{t['cat']}</td>\n"
            f"          <td>{t['factors']}</td>\n"
            f"          <td class=\"num\">{t['explain']}</td>\n"
            f"          <td>{t['status']}</td>\n"
            "        </tr>"
        )
    block = "\n".join(rows)
    pattern = (
        r"(<h3>传统核保因子（核心，持续挖深）</h3>\s*<table>[\s\S]*?<tbody>\s*)"
        r"([\s\S]*?)(\s*</tbody>)"
    )
    m = re.search(pattern, html)
    if not m:
        raise ValueError("未找到传统核保因子表格")
    return html[: m.start(2)] + block + html[m.end(2) :]


def replace_adas_table(html: str, adas: list[dict]) -> str:
    rows = []
    for i, a in enumerate(adas):
        cls = ' class="highlight"' if i < 4 and a["maturity"] != "中长期布局" else ""
        if "远期" in a["dir"] or "中长期" in a["maturity"]:
            cls = ""
        else:
            cls = ' class="highlight"'
        rows.append(
            f"<tr{cls}>\n"
            f"          <td>{a['dir']}</td>\n"
            f"          <td>{a['obs']}</td>\n"
            f"          <td>{a['risk']}</td>\n"
            f"          <td>{a['maturity']}</td>\n"
            "        </tr>"
        )
    block = "\n".join(rows)
    pattern = (
        r"(<h3>智驾类因子（占比小、斜率大 —— 应系统投入）</h3>\s*<table>[\s\S]*?<tbody>\s*)"
        r"([\s\S]*?)(\s*</tbody>)"
    )
    m = re.search(pattern, html)
    if not m:
        raise ValueError("未找到智驾类因子表格")
    return html[: m.start(2)] + block + html[m.end(2) :]


def replace_todos(html: str, todos: list[dict]) -> str:
    items = []
    for t in todos:
        extra = ""
        if t["conclusion"]:
            extra = f"<br><span style=\"color:#0e7c7b\">结论：{t['conclusion']}</span>（{t['status']}）"
        elif t["status"] and t["status"] != "待讨论":
            extra = f"<br><span style=\"color:#0e7c7b\">状态：{t['status']}</span>"
        items.append(
            "<li>\n"
            f"        <div class=\"todo-num\">{t['n']}</div>\n"
            "        <div>\n"
            f"          <b>{t['topic']}</b>"
            f"{extra}\n"
            "        </div>\n"
            "      </li>"
        )
    block = "\n".join(items)
    pattern = r'(<ul class="todo-list">\s*)([\s\S]*?)(\s*</ul>)'
    m = re.search(pattern, html)
    if not m:
        raise ValueError("未找到待拍板清单")
    return html[: m.start(2)] + block + html[m.end(2) :]


def patch_meta(html: str, meta: dict) -> str:
    if meta.get("doc_badge"):
        html = re.sub(
            r'(<div class="badge">)(.*?)(</div>)',
            rf'\1{meta["doc_badge"]}\3',
            html,
            count=1,
        )
    if meta.get("doc_title") and meta.get("doc_subtitle"):
        html = re.sub(
            r"(<h1>)([\s\S]*?)(</h1>)",
            rf'\1{meta["doc_title"]}<br><em>{meta["doc_subtitle"]}</em>\3',
            html,
            count=1,
        )
    if meta.get("doc_lead"):
        html = re.sub(
            r'(<p class="lead">\s*)([\s\S]*?)(\s*</p>)',
            rf'\1{meta["doc_lead"]}\3',
            html,
            count=1,
            flags=re.S,
        )
    if meta.get("doc_date"):
        html = re.sub(
            r'(<span class="chip">)(\d{4}-\d{2}-\d{2})(</span>)',
            rf'\1{meta["doc_date"]}\3',
            html,
            count=1,
        )
    if meta.get("doc_owner"):
        html = re.sub(
            r'(<span class="chip strong">)(.*?)(</span>)',
            rf'\1{meta["doc_owner"]}\3',
            html,
            count=1,
        )
    if meta.get("leader_point_1") and meta.get("leader_point_2"):
        callout = (
            f"<strong>（1）明确整体风险研究框架：</strong>\n"
            f'      {meta["leader_point_1"]}\n'
            f"      <br><br>\n"
            f"      <strong>（2）新因子积累不足：</strong>\n"
            f'      {meta["leader_point_2"]}'
        )
        html = re.sub(
            r'(<div class="callout">\s*)([\s\S]*?)(\s*</div>\s*<h3>本稿要回答的三件事</h3>)',
            rf"\1{callout}\3",
            html,
            count=1,
        )
    if meta.get("口径提醒"):
        html = re.sub(
            r'(口径提醒：下文「人相关 80%–90%」是<strong>因果/损失归因口径</strong>（事故主因多在驾驶行为）；\s*'
            r"筛查模型里「车龄、车价」贡献度可以很高——那是<strong>可观测代理变量</strong>，不等于「车本身造成了多数风险」。讨论时务必把两套口径拆开讲。)",
            meta["口径提醒"],
            html,
            count=1,
        )
    if meta.get("分层建议结论"):
        html = re.sub(
            r"(<strong>建议结论（供拍板）：</strong>\s*)([\s\S]*?)(\s*</div>)",
            rf'\1{meta["分层建议结论"]}\3',
            html,
            count=1,
        )
    return html


def sync(xlsx: Path, html_path: Path, also_copy: Path | None = None) -> None:
    data = read_workbook(xlsx)
    html = html_path.read_text(encoding="utf-8")

    seg_js = emit_segments_js(data["segments"])
    fac_js = emit_factors_js(data["factors"])

    html2, n1 = re.subn(
        r"  const SEGMENTS = \[[\s\S]*?\n  \];",
        seg_js,
        html,
        count=1,
    )
    if n1 != 1:
        raise ValueError("替换 SEGMENTS 失败")
    html3, n2 = re.subn(
        r"  const FACTORS = \{[\s\S]*?\n  \};",
        fac_js,
        html2,
        count=1,
    )
    if n2 != 1:
        raise ValueError("替换 FACTORS 失败")

    html3 = replace_trad_table(html3, data["trad"])
    html3 = replace_adas_table(html3, data["adas"])
    html3 = replace_todos(html3, data["todos"])
    html3 = patch_meta(html3, data["meta"])

    html_path.write_text(html3, encoding="utf-8")
    if also_copy:
        also_copy.write_text(html3, encoding="utf-8")
    print(f"✅ synced from {xlsx.name} → {html_path}", file=sys.stderr)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--html", type=Path, default=DEFAULT_HTML)
    ap.add_argument("--also-copy", type=Path, default=DEFAULT_HTML_COPY)
    args = ap.parse_args()
    if not args.xlsx.exists():
        raise SystemExit(f"xlsx not found: {args.xlsx}")
    if not args.html.exists():
        raise SystemExit(f"html not found: {args.html}")
    sync(args.xlsx, args.html, args.also_copy)


if __name__ == "__main__":
    main()
