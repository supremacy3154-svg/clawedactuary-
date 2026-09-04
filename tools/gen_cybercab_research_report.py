#!/usr/bin/env python3
"""Generate an A4 research-note .docx for the Tesla Cybercab memo (not site HTML)."""
from __future__ import annotations

import argparse
import smtplib
import ssl
import sys
from email import encoders
from email.header import Header
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from smtp_config import FROM_ADDR, PASSWORD, SMTP_HOST, SMTP_PORT, USERNAME

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data" / "reports"
OUT_NAME = "2026-09-04-Tesla-Cybercab-机动车身份与保险监管口径-研究报告.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x1A, 0x18, 0x14)
MUTED = RGBColor(0x5A, 0x57, 0x50)
RULE = "1F4E79"
BODY_SIZE = 12
EAST = "宋体"
EAST_HEAD = "黑体"


def set_run_font(run, *, east: str = EAST, size: Pt | None = None, bold: bool | None = None, color=None):
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), east)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "8FAADC")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_bottom_border(paragraph, color: str = RULE, sz: str = "12") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_paragraph_spacing(p, *, before=0, after=8, line=22, first_line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)


def add_body(doc, text: str, *, first_line=0.74):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=8, line=22, first_line=first_line)
    run = p.add_run(text)
    set_run_font(run, size=Pt(BODY_SIZE), color=INK)
    return p


def add_heading_cn(doc, text: str, level: int):
    p = doc.add_paragraph()
    if level == 1:
        set_paragraph_spacing(p, before=16, after=10, line=24)
        run = p.add_run(text)
        set_run_font(run, east=EAST_HEAD, size=Pt(16), bold=True, color=NAVY)
        add_bottom_border(p, sz="8")
    else:
        set_paragraph_spacing(p, before=12, after=6, line=22)
        run = p.add_run(text)
        set_run_font(run, east=EAST_HEAD, size=Pt(13), bold=True, color=NAVY)
    return p


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=12, line=18)
    run = p.add_run(text)
    set_run_font(run, size=Pt(10.5), color=MUTED)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, east=EAST_HEAD, size=Pt(10.5), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1F4E79")
        set_cell_border(cell)
    for r_i, row in enumerate(rows):
        fill = "F2F6FB" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=Pt(10.5), color=INK)
            shade_cell(cell, fill)
            set_cell_border(cell)
    doc.add_paragraph()
    return table


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("专题研究报告  ·  Tesla Cybercab 机动车身份与保险监管口径")
    set_run_font(run, size=Pt(9), color=MUTED)
    add_bottom_border(hp, color="8FAADC", sz="6")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("仅供研究参阅，不构成法律、投保或投资建议    第 ")
    set_run_font(run, size=Pt(9), color=MUTED)

    def add_field(paragraph, instr: str):
        r1 = paragraph.add_run()
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        r1._r.append(fld1)
        r2 = paragraph.add_run()
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = instr
        r2._r.append(t)
        r3 = paragraph.add_run()
        fld3 = OxmlElement("w:fldChar")
        fld3.set(qn("w:fldCharType"), "end")
        r3._r.append(fld3)

    add_field(fp, " PAGE ")
    run = fp.add_run(" 页")
    set_run_font(run, size=Pt(9), color=MUTED)


def build_document(path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    add_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST)

    for _ in range(3):
        doc.add_paragraph()
    k = doc.add_paragraph()
    k.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = k.add_run("专题研究报告")
    set_run_font(run, east=EAST_HEAD, size=Pt(14), bold=True, color=NAVY)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bottom_border(line, sz="18")

    t1 = doc.add_paragraph()
    t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(t1, before=18, after=6, line=32)
    run = t1.add_run("Tesla Cybercab 在美国的机动车法律身份")
    set_run_font(run, east=EAST_HEAD, size=Pt(22), bold=True, color=INK)

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(t2, before=0, after=12, line=32)
    run = t2.add_run("与相关监管要求的保险类型")
    set_run_font(run, east=EAST_HEAD, size=Pt(22), bold=True, color=INK)

    en = doc.add_paragraph()
    en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(en, before=4, after=18, line=20)
    run = en.add_run("Vehicle Classification and Statutory Insurance Requirements\nfor Tesla Cybercab in the United States")
    set_run_font(run, size=Pt(11), color=MUTED)

    meta_rows = [
        ("报告类型", "海外车险监管专题研究"),
        ("研究问题", "Cybercab 在美国联邦法上是否属于机动车；监管要求投保哪一类保险"),
        ("资料范围", "美国联邦与州法律原文，以及英文公开报道"),
        ("截稿日期", "2026年9月4日"),
        ("版本", "第2稿，已删配图，并按书面研究报告改写正文"),
    ]
    table = doc.add_table(rows=len(meta_rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k_txt, v_txt) in enumerate(meta_rows):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k_txt)
        set_run_font(r0, east=EAST_HEAD, size=Pt(11), bold=True, color=NAVY)
        shade_cell(c0, "E7EEF7")
        set_cell_border(c0)
        c1.text = ""
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v_txt)
        set_run_font(r1, size=Pt(11), color=INK)
        set_cell_border(c1)
        c0.width = Cm(3.4)
        c1.width = Cm(12.2)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note, before=28, after=0, line=18)
    run = note.add_run("本稿供研究参阅，不构成法律意见、投保建议、承保决策或投资建议。\n文中判断依据公开材料整理，不代表任何机构观点。")
    set_run_font(run, size=Pt(10), color=MUTED)

    doc.add_page_break()

    add_heading_cn(doc, "摘要", 1)
    add_body(
        doc,
        "Tesla 已于 2026 年第一季度在得克萨斯超级工厂投产 Cybercab。该车是两座 Robotaxi，不设方向盘和踏板。"
        "9 月 3 日奥斯汀发布会前后，Tesla Robotaxi, LLC 已将 VIN 前缀 5YJA 的车辆登记进入德州机动车承运人资质系统。"
        "公开查询显示，Cybercab 一度登记 45 台，与 Model Y 合计授权约 420 台。",
    )
    add_body(
        doc,
        "本报告讨论两个问题：Cybercab 在美国联邦法上是否属于机动车；如果属于，相关监管要求投保哪一类保险，限额如何确定。",
    )
    add_body(
        doc,
        "先看分类。按照 49 U.S.C. §30102，车辆由机械动力驱动、主要为公路使用而制造，即属机动车。"
        "条文没有把方向盘、踏板或人类驾驶人规定为必备条件。"
        "NHTSA 2016 年解释函和 2022 年乘员保护终规，都把没有传统操控件的自动驾驶车辆继续纳入机动车安全标准。"
        "Tesla 称 Cybercab 已按联邦机动车安全标准自我认证，因此不受每 12 个月 2,500 辆临时豁免上限约束。",
    )
    add_body(
        doc,
        "再看保险。联邦安全法不规定保单名称，也不规定责任限额。投保义务写在各州财务责任法中，测试、部署和载客会再叠加一层要求。"
        "德州目前要求投保机动车责任险，法律同时承认自保。按网约载客运营时，自动化机动车视同运输网络公司司机："
        "待单阶段适用分项限额，乘客上车后适用 100 万美元综合限额。"
        "若在加州开展无人驾驶载客，制造商测试许可和公用事业委员会客运许可还须各自备妥 500 万美元财务责任。"
        "Tesla 尚未公布 Cybercab 保单。私人乘用车保险不能覆盖车队网约。",
    )
    add_body(
        doc,
        "核保时需要确认车辆当时处于待单、载客还是测试，以及持有哪一类许可。"
        "分项限额、综合限额和制造商财务责任口径不同，数字不宜直接比较。",
    )
    kw = doc.add_paragraph()
    set_paragraph_spacing(kw, before=6, after=10, line=20, first_line=0)
    run = kw.add_run("关键词：")
    set_run_font(run, east=EAST_HEAD, size=Pt(12), bold=True, color=NAVY)
    run = kw.add_run("Tesla Cybercab；机动车；FMVSS 自我认证；运输网络公司责任险；SB 2807；加州车辆法典第 38750 条")
    set_run_font(run, size=Pt(12), color=INK)

    add_heading_cn(doc, "目录", 1)
    toc_items = [
        "一、研究背景与问题",
        "二、联邦法上的机动车认定",
        "三、自我认证与豁免上限",
        "四、保险义务怎样分层",
        "五、德克萨斯州：当前部署地",
        "六、加利福尼亚州对照",
        "七、产品责任与尚未公开的保单",
        "八、对国内财产险承保的启示",
        "九、后续跟踪事项",
        "附录一  主要英文来源",
        "附录二  局限与口径说明",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=2, after=2, line=22, first_line=0)
        run = p.add_run(item)
        set_run_font(run, size=Pt(12), color=INK)

    add_heading_cn(doc, "一、研究背景与问题", 1)
    add_body(
        doc,
        "Cybercab 与常见乘用车不同：两座，不设方向盘和踏板，专门用于 Robotaxi。"
        "强制责任险是否适用，首先要确认它是否属于机动车；限额落在哪一档，再看实际用途。这两项判断需要分开处理。",
    )
    add_body(
        doc,
        "本报告回答两个问题。第一，这台车在美国联邦法上是否属于机动车。"
        "第二，如果属于，相关监管要求的是哪一类保险，限额如何确定。",
    )
    add_body(
        doc,
        "材料来自美国联邦和州的法律原文，以及英文公开报道。"
        "生产进度和自我认证的表述，转述 Electrek 对 Tesla 2026 年第一季度财报电话会的报道。"
        "德州名册中的台数，转述 TeslaNorth、Teslarati 对公开查询的报道。"
        "报告未获得保单样本，也没有出险数据。事实判断不以中文媒体为依据。",
    )

    add_heading_cn(doc, "二、联邦法上的机动车认定", 1)
    add_body(
        doc,
        "联邦法把机动车定义写在 49 U.S.C. §30102(a)(7)：由机械动力驱动或牵引，主要为公共街道和公路制造；只在轨道上运行的除外。"
        "Cybercab 为公路载客设计，符合这一定义。"
        "方向盘、踏板、人类驾驶人，条文均未列为必备条件。",
    )
    add_body(
        doc,
        "2016 年 2 月 4 日，国家公路交通安全管理局就 Google 自动驾驶项目发出解释函，讨论去掉方向盘和踏板之后安全标准如何适用。"
        "信中把自动驾驶系统视为机动车设备；在 Google 那套设计里，驾驶人被解释为这套系统。"
        "信同时写明：车辆仍须按当时有效的联邦机动车安全标准自我认证。"
        "解释函不能改写需要修规或豁免才能调整的条文。"
        "当时不少标准仍按脚刹车、转向管柱来写，认证方法因此变得困难，车辆的机动车身份并不因此改变。",
    )
    add_body(
        doc,
        "2022 年 3 月 30 日，《联邦公报》87 FR 18560 发布《配备自动驾驶系统车辆的乘员保护》终规，同年 9 月 26 日生效。"
        "终规调整了 200 系列碰撞防护标准中的用语，使「驾驶人座位」「方向盘」等表述也能适用于没有传统操控件的车辆，"
        "并要求这类车辆仍然提供与现行乘用车相当的乘员保护。适用对象仍是机动车中的乘用车。",
    )
    add_caption(doc, "表1  联邦认定依据")
    add_table(
        doc,
        ["规范", "含义"],
        [
            ["49 U.S.C. §30102(a)(7)", "公路用机械动力车辆即机动车，条文未要求方向盘或踏板"],
            ["NHTSA 2016年2月4日解释函", "自动驾驶系统属机动车设备；车上没有人类驾驶人，车辆仍须认证"],
            ["87 FR 18560，2022年3月30日", "没有传统操控件的自动驾驶车辆，仍须提供同等碰撞保护"],
            ["49 CFR Part 555", "只有申请临时豁免时，才受每12个月2,500辆上限约束"],
        ],
    )
    add_body(
        doc,
        "认定顺序是清楚的。先确认它属于机动车，再讨论如何证明符合安全标准。"
        "保险义务跟在分类之后，限额改写不了车辆的法律身份。",
    )

    add_heading_cn(doc, "三、自我认证与豁免上限", 1)
    add_body(
        doc,
        "美国不对新车做事前型式批准。制造商自行声明：出厂当日符合全部适用的联邦机动车安全标准，发现缺陷后须召回。"
        "某一条标准确实无法满足时，可按 49 U.S.C. §30113 和 49 CFR Part 555 申请临时豁免。"
        "除经济困难等少数情形外，申请时必须声明：获豁免车辆在任意 12 个月里，在美国销售不超过 2,500 辆。"
        "这一上限约束的是豁免车辆。已经自我认证、声称符合全部标准的车辆，不受该上限约束。",
    )
    add_body(
        doc,
        "Electrek 报道 Tesla 2026 年第一季度电话会：马斯克确认 Cybercab 已经投产；"
        "车辆工程副总裁 Lars Moravy 在社交网络上用一个 No 回答 2,500 辆上限是否适用于 Cybercab。"
        "报道还提到，工厂画面中的车辆已经贴上联邦合规标签。这些内容均来自制造商自身表述。"
        "美国安全法不对单一车型发放符合性批准，国家公路交通安全管理局也没有发布过 Cybercab 专属批准函。"
        "此后如果检测或缺陷调查认定不符合标准，监管机构仍可执法并要求召回。",
    )
    add_body(
        doc,
        "德州把同一套分类写进了州法。参议院 SB 2807 在运输法典中增设自动化机动车专章，自 2026 年 5 月 28 日起，"
        "商业无人驾驶载客、载货均须执行：车上的自动驾驶系统须符合联邦法律，包括联邦机动车安全标准，运营还须取得德州机动车管理局授权。"
        "TeslaNorth 和 Teslarati 查阅公开记录后写到：Tesla Robotaxi, LLC 自 2026 年 8 月 31 日起登记 2026 款 Cybercab，"
        "VIN 前缀 5YJA，与 Model Y 的 7SAYG 不是同一系列。"
        "当天从 7 台增加到 45 台，到发布会当天授权合计大约 420 台。"
        "进入名册只说明已取得商业运营授权。乘客能否在 Robotaxi 应用中叫到这款车，还需要另行核实。",
    )

    add_heading_cn(doc, "四、保险义务怎样分层", 1)
    add_body(
        doc,
        "联邦《机动车安全法》规范新车性能和缺陷，不规定保单名称，也不规定责任限额。"
        "美国没有一张全国通用的 Cybercab 自动驾驶保单。"
        "实际需要满足的，是三层义务：上路财务责任，自动驾驶测试或部署许可，以及网约或包车客运许可。",
    )
    add_caption(doc, "表2  按场景划分的法定保险下限")
    add_table(
        doc,
        ["场景", "法定口径", "下限"],
        [
            ["德州普通上路", "运输法典第 601.072 条", "人身伤害每人 3 万美元、每事故 6 万美元，财产损失 2.5 万美元"],
            ["德州网约待单", "保险法第 1954.052 条", "人身伤害每人 5 万美元、每事故 10 万美元，财产损失 2.5 万美元"],
            ["德州预约定程", "保险法第 1954.053 条", "死亡、人身伤害与财产损失综合限额 100 万美元"],
            ["加州车主财务责任", "车辆法典第 16056 条，2025 年起保单", "人身伤害每人 3 万美元、每事故 6 万美元，财产损失 1.5 万美元"],
            ["加州制造商测试或部署", "车辆法典第 38750 条及 13 CCR §227.04", "500 万美元保险、保证或自保"],
            ["加州自动驾驶客运项目", "公用事业委员会相关决定", "公共责任与财产损失 500 万美元；有雇员另须工伤补偿"],
        ],
    )
    add_body(
        doc,
        "表中所列均为法定下限，口径并不相同。德州前两档将人身伤害和财产损失分开规定，预约定程则采用综合限额。"
        "加州的 500 万美元，是制造商或客运许可要求的财务责任证明，形式包括保险、保证和自保。"
        "产品责任不在这些下限之内。",
    )
    add_body(
        doc,
        "法律规定应当投保的险种，与市场上实际签发的保单，需要分开看待。"
        "车队运营方购买的，多半是商用车险责任保障，再按网约或客运许可提高限额。"
        "私人乘用车保单把网约、出租和车队商业使用列为除外。"
        "Tesla Insurance 在华盛顿等地备案的，是私人乘用车实时定价产品，不能充当 Cybercab 车队的法定责任险。",
    )

    add_heading_cn(doc, "五、德克萨斯州：当前部署地", 1)
    add_body(
        doc,
        "SB 2807 所称自动化机动车，是指装有自动驾驶系统的机动车。"
        "运输法典第 545.455 条规定，系统激活期间，车上必须具有机动车责任保险或自保，额度不得低于本州或联邦法律对该车类型和用途规定的金额。"
        "使用无人驾驶车辆从事商业载客，还须向德州机动车管理局申请授权，并提交应急互动方案。"
        "申请材料中须书面确认保险已经落实。"
        "法律没有另外规定一个 500 万美元的自动驾驶险种，限额随车辆类型和用途确定。"
        "部分英文二手报道把加州的 500 万美元直接套用到德州，与法案文本不符。",
    )
    add_body(
        doc,
        "同一部法案在保险法典中增加第 1954.003 条：运输法典中的自动化机动车，在第 1954 章 B 分章中视为运输网络公司司机，该分章的承保要求一并适用。"
        "已登录数字网络、可以接单但尚未进入预约定程时，按第 1954.052 条，人身伤害每人至少 5 万美元、每事故 10 万美元，财产损失 2.5 万美元。"
        "进入预约定程后，按第 1954.053 条，死亡、人身伤害和财产损失的综合限额至少 100 万美元，并在本州要求的范围内提供未投保驾驶人和人身伤害保护。"
        "职业法典第 2402 章也将通过数字网络安排的自动化机动车行程纳入运输网络公司规则。",
    )
    add_body(
        doc,
        "系统激活时，交通和机动车法律上的操作者是这套系统。违法通知可以向车辆所有人或授权持有人送达。"
        "公开名册上的持有人是 Tesla Robotaxi, LLC。"
        "对保险公司而言，被保险人应当是运营公司和登记所有人。驾驶座上有没有人，不改变承保对象。"
        "私人乘用车保单中的合同载客除外一旦触发，赔付就会落到车队商用车险和运输网络公司责任保障上。",
    )
    add_body(
        doc,
        "名册上有车，就要维持保险。维持保险的义务，与乘客已经能够在应用中叫到 Cybercab，是两件需要分开核实的事实。"
        "TeslaNorth 写过：授权清单是监管上限，路上实际在跑、能够派单的数量需要另行核对。"
        "核保时应当核对授权持有人、VIN 清单、这款车有没有进入预约定程，以及出险时是待单还是载客。",
    )

    add_heading_cn(doc, "六、加利福尼亚州对照", 1)
    add_body(
        doc,
        "加州把自动驾驶汽车的测试和部署写在车辆法典第 16.6 编。"
        "第 38750 条规定，制造商要在本州公路上开始测试，须先备妥 500 万美元的保险、保证或自保，并按机动车管理局的要求提交。"
        "部署许可同样须维持这一金额。"
        "加州法规汇编第 13 编第 227.04 条把证明形式写明为：加州持牌保险人出具的保单、持牌或合格溢额保证，或者自保证书。"
        "选择自保的，还须提交近三年经审计、净资产不低于 500 万美元的报表。"
        "这 500 万美元约束的是制造商对人身伤害、死亡和财产损失判决的偿付能力，与车主按第 16056 条承担的普通财务责任相互独立。",
    )
    add_body(
        doc,
        "向乘客收费，还须通过加州公用事业委员会。"
        "委员会的自动驾驶客运试点和一期部署项目，要求公共责任和财产损失保险 500 万美元，高于包车承运人通则第 115 号系列的普通限额。"
        "有雇员的，还须有工伤补偿。保险人须向委员会电子备案，证明保单持续有效。",
    )
    add_body(
        doc,
        "2026 年 3 月，委员会负责消费者政策、交通和执法的副执行主任 Pat Tsen 在访谈中说："
        "Tesla 在加州的约车，按有人驾驶的包车客运来管，没有进入自动驾驶客运项目。"
        "加州把自动驾驶定为驾驶自动化 3 级及以上，Tesla 的系统按 2 级管理。"
        "公司拿到的是包车承运许可，和礼宾车公司是同一类；坐在驾驶座上的人，就是驾驶人。"
        "Electrek 在 2026 年 3 月 25 日报道了这次谈话。"
        "TechCrunch 在 2025 年许可获批时已经写过：这张包车许可不覆盖自动驾驶测试或部署，"
        "Tesla 当时没有申请委员会的自动驾驶客运项目，也没有加州机动车管理局的无人驾驶载客许可。",
    )
    add_body(
        doc,
        "此后如果要把没有方向盘的 Cybercab 在加州向公众提供无人驾驶行程，"
        "机动车管理局的自动驾驶许可和委员会的自动驾驶客运许可都需要办理。"
        "保险也会从包车通则，提高到制造商 500 万美元再加客运 500 万美元。"
        "这两张许可办妥之前，加州现在的 Tesla 约车仍按有人驾驶的包车责任险管理。",
    )

    add_heading_cn(doc, "七、产品责任与尚未公开的保单", 1)
    add_body(
        doc,
        "前面这些条文要求的，都是机动车第三者责任，以及测试、网约、客运许可中加高的限额。证明形式包括保险、保证和自保。"
        "软件出错、最小风险策略选错、远程坐席指挥失误，仍可能构成产品责任或专业责任索赔。"
        "联邦和德州、加州的财务责任法，都没有把产品责任险写成上路的前提。"
        "出了事故，索赔先找机动车责任险；限额不够，或者争点落在设计缺陷上，再转到产品责任。",
    )
    add_body(
        doc,
        "Tesla 没有公布 Cybercab 的保单名称、承保公司、是否自保，也没有说明限额有没有高于法定下限。"
        "马斯克多次把 Robotaxi 的运营成本目标说成每英里大约 0.20 美元，里面含保险。"
        "这是公司内部的成本假设，替代不了法定保单。"
        "核保时需要看到的材料包括：Tesla Robotaxi, LLC 的商用车险责任限额，有没有批注运输网络公司分阶段保障，"
        "有没有超额责任，远程协助人员是否纳入工伤补偿，产品责任是否由制造商保单来接。",
    )

    add_heading_cn(doc, "八、对国内财产险承保的启示", 1)
    add_body(
        doc,
        "对国内财产险公司来说，有三点可以直接对照。"
        "第一，先确认车辆是否属于机动车，再讨论保险；没有方向盘，并不能免除强制责任险。"
        "第二，同一辆车在待单、载客、测试、跨州时，限额并不相同，保单需要按行程状态切换，一张「智驾险」覆盖不了全部状态。"
        "第三，制造商 500 万美元和客运 500 万美元都是许可门槛，回答不了损失有多高；定价仍须回到车队里程、运营设计域和软件版本。",
    )

    add_heading_cn(doc, "九、后续跟踪事项", 1)
    items = [
        "自我认证能否经得起检测和缺陷调查。国家公路交通安全管理局如果认定碰撞防护、灯光或制动标准未满足，会启动缺陷调查和召回。每 12 个月 2,500 辆的豁免上限不会自动适用于已经自我认证的车辆，但认证主张可以被推翻。",
        "德州名册何时转为真正能够派单的运力。需要观察 Tesla Robotaxi 应用是否派发 5YJA 前缀车辆，以及出险后是否按第 1954 章分阶段限额赔付。",
        "加州许可是否升级。没有方向盘的载客，需要机动车管理局自动驾驶许可和委员会自动驾驶项目；一旦申请，500 万美元财务责任就会成为可以核对的备案材料。",
        "实际保单是否高于法定下限。100 万美元和 500 万美元都是许可门槛。严重伤亡事故的判决可能高于门槛，超额责任和再保安排目前未见披露。",
        "产品责任和机动车责任如何分单。感知误判、远程坐席指令和车辆机械故障如果由不同保单承接，追偿顺序会决定案均赔付落在哪一份合同上。",
    ]
    for i, text in enumerate(items, 1):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=6, line=22, first_line=0)
        run = p.add_run(f"{i}. {text}")
        set_run_font(run, size=Pt(BODY_SIZE), color=INK)

    add_heading_cn(doc, "附录一  主要英文来源", 1)
    add_body(
        doc,
        "下列来源均为英文监管原文或英文公开报道。中文为作者转写。",
        first_line=0,
    )
    refs = [
        "49 U.S.C. §30102. Definitions. https://uscode.house.gov/",
        "NHTSA, Interpretation letter to Google, Inc., Feb. 4, 2016. https://www.nhtsa.gov/interpretations/google-compiled-response-12-nov-15-interp-request-4-feb-16-final",
        "Occupant Protection for Vehicles With Automated Driving Systems, 87 Fed. Reg. 18560 (Mar. 30, 2022). https://www.govinfo.gov/content/pkg/FR-2022-03-30/html/2022-05426.htm",
        "49 CFR Part 555, Temporary Exemption from Motor Vehicle Safety and Bumper Standards. https://www.ecfr.gov/current/title-49/part-555",
        "NHTSA interpretation GF005146 (Part 555 2,500-vehicle cap; no carry-forward).",
        "Texas S.B. 2807, 89th Leg., R.S. (enrolled). https://capitol.texas.gov/tlodocs/89R/billtext/html/SB02807F.HTM",
        "Tex. Transp. Code §545.455, §601.072.",
        "Tex. Ins. Code §§1954.003, 1954.052, 1954.053.",
        "Cal. Veh. Code §§16056, 38750; Cal. Code Regs. tit. 13, §227.04.",
        "CPUC, Autonomous Vehicle Program Applications—General Guidance (Aug. 2024); Decision 18-05-043; Decision 20-11-046.",
        "Electrek, “Tesla confirms Cybercab production has started despite delays in unsupervised driving,” Apr. 23, 2026.",
        "Electrek, “California regulator confirms Tesla is not operating an autonomous vehicle service,” Mar. 25, 2026.",
        "TechCrunch, “What Tesla can and can’t do in California with its new passenger transportation permit,” Mar. 18, 2025.",
        "TeslaNorth, “Tesla Cybercab Just Made Texas’s Official Robotaxi Roster,” Sept. 1, 2026; “Tesla Robotaxi Roster Hits 420 Vehicles,” Sept. 2, 2026.",
        "Teslarati, “Tesla Cybercab fleet grows in Austin ahead of launch event.”",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=4, line=18, first_line=0)
        run = p.add_run(f"[{i}]  {ref}")
        set_run_font(run, size=Pt(10.5), color=INK)

    add_heading_cn(doc, "附录二  局限与口径说明", 1)
    limits = [
        "联邦定义与安全标准以 49 U.S.C. 第 301 章、NHTSA 2016 年 2 月 4 日解释函、87 FR 18560 及 49 CFR Part 555 为准。Tesla 自我认证是制造商主张，监管机构未发布 Cybercab 专属批准。",
        "德州保险与授权条文以 SB 2807 入法文本、运输法典第 545.455 条、保险法第 1954 章为准。TxMCCS 台数转述 TeslaNorth 与 Teslarati 对公开查询的报道，名册会变动，截稿日为 2026 年 9 月 4 日。",
        "加州 500 万美元要求引自车辆法典第 38750 条、13 CCR §227.04 及公用事业委员会自动驾驶项目指引。Tsen 谈话转述 Electrek 2026 年 3 月 25 日报道。Tesla 现行加州约车许可可能在截稿后变更。",
        "表 2 金额为法定下限。分项限额与综合限额口径不同，不宜直接比较。",
        "关于保单形态、产品责任分单与国内迁移的表述，是作者判断，不是企业披露。",
    ]
    for text in limits:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=6, line=20, first_line=0)
        run = p.add_run("·  " + text)
        set_run_font(run, size=Pt(10.5), color=INK)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def send_docx(to: str, path: Path, subject: str, body: str) -> None:
    msg = MIMEMultipart()
    msg["From"] = FROM_ADDR
    msg["To"] = to
    msg["Subject"] = Header(subject, "utf-8")
    msg.attach(MIMEText(body, "plain", "utf-8"))
    part = MIMEBase("application", "vnd.openxmlformats-officedocument.wordprocessingml.document")
    part.set_payload(path.read_bytes())
    encoders.encode_base64(part)
    ascii_name = "2026-09-04-Tesla-Cybercab-US-MV-Insurance-Research-Note-v2.docx"
    part.add_header("Content-Disposition", "attachment", filename=ascii_name)
    msg.attach(part)
    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ctx) as s:
        s.login(USERNAME, PASSWORD)
        refused = s.sendmail(FROM_ADDR, [to], msg.as_bytes())
        if refused:
            raise smtplib.SMTPRecipientsRefused(refused)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--to", default="yanghailin508@pingan.com.cn")
    parser.add_argument("--send", action="store_true")
    args = parser.parse_args()
    out = OUT_DIR / OUT_NAME
    path = build_document(out)
    print(f"wrote {path} ({path.stat().st_size} bytes)", file=sys.stderr)
    if args.send:
        send_docx(
            args.to,
            path,
            "【研究报告·修订】Tesla Cybercab 在美国的机动车身份与保险监管口径",
            "附件为修订后的 Word 研究报告：已去掉两张生成示意图，并按书面研究报告语气改写了中文。\n\n"
            "结构仍是封面、摘要、目录、九章正文和英文来源附录，保留表1、表2。\n"
            "不构成法律、投保或投资建议。",
        )
        print(f"sent to {args.to}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
